#!/usr/bin/env python3
import argparse
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

from jira_common import JiraClient, adf_text, add_env_arg, print_json
from jira_scrum_workbook_import import module_for_pbi, priority_name, slug, transition_if_needed


def clean(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip()


def load_rows(path: str, sheet_name: str) -> List[Dict[str, str]]:
    wb = load_workbook(filename=Path(path).expanduser(), data_only=True)
    ws = wb[sheet_name]
    headers = [clean(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        item = {headers[idx]: clean(row[idx]) if idx < len(headers) else '' for idx in range(len(headers)) if headers[idx]}
        rows.append(item)
    return rows


def build_plan(file_path: str, project_key: str) -> Dict[str, Any]:
    user_stories = load_rows(file_path, 'User Stories')
    product_backlog = load_rows(file_path, 'Product Backlog')
    sprint_backlog = load_rows(file_path, 'Sprint Backlog')
    pbi_map = {row.get('ID', ''): row for row in product_backlog}

    tasks = []
    for idx, row in enumerate(sprint_backlog, start=2):
        pbi_id = row.get('PBI ID', '')
        pbi = pbi_map.get(pbi_id, {})
        sprint = row.get('Sprint', '')
        sprint_goal = row.get('Sprint Goal', '')
        task = row.get('Task', '')
        pic = row.get('PIC', '')
        est = row.get('Estimasi (Jam)', '')
        status = row.get('Status', '')
        dod = row.get('Catatan/DoD', '')
        pbi_summary = pbi.get('User Story / PBI', '')
        pbi_priority = priority_name(pbi.get('Prioritas', ''))
        module = module_for_pbi(pbi_summary, user_stories) if pbi_summary else 'General'
        labels = [
            'import-simasu',
            'sprint-task',
            f'src-sb-row-{idx}',
            f'sprint-{slug(sprint)}' if sprint else 'sprint-unknown',
            f'pbi-{slug(pbi_id)}' if pbi_id else 'pbi-unknown',
            f'module-{slug(module)}' if module else 'module-general',
        ]
        if pic:
            labels.append(f'pic-{slug(pic)}')
        description_lines = [
            f'Source Row: {idx}',
            f'Sprint: {sprint}',
            f'Sprint Goal: {sprint_goal}',
            f'PBI ID: {pbi_id}',
            f'Related PBI: {pbi_summary}',
            f'Module: {module}',
            f'PIC: {pic}',
            f'Estimate Hours: {est}',
            f'Definition of Done: {dod}',
        ]
        tasks.append({
            'source_id': f'SB-ROW-{idx}',
            'summary': task,
            'status': status,
            'priority': pbi_priority,
            'labels': labels,
            'description': adf_text('\n'.join(description_lines)),
        })

    return {
        'project_key': project_key,
        'count': len(tasks),
        'tasks': tasks,
    }


def existing_map(client: JiraClient, project_key: str) -> Dict[str, str]:
    result = client.search(f'project = {project_key} AND labels = "sprint-task"', limit=200, fields='labels,summary,status')
    mapping = {}
    for issue in result.get('issues', []):
        labels = issue.get('fields', {}).get('labels', []) or []
        for label in labels:
            if label.startswith('src-sb-row-'):
                mapping[label] = issue['key']
    return mapping


def create_task(client: JiraClient, project_key: str, task: Dict[str, Any]) -> str:
    fields: Dict[str, Any] = {
        'project': {'key': project_key},
        'issuetype': {'name': 'Task'},
        'summary': task['summary'],
        'description': task['description'],
        'labels': task['labels'],
    }
    if task.get('priority'):
        fields['priority'] = {'name': task['priority']}
    result = client.issue_create({'fields': fields})
    return result['key']


def apply_plan(client: JiraClient, plan: Dict[str, Any]) -> Dict[str, Any]:
    project_key = plan['project_key']
    existing = existing_map(client, project_key)
    created = []
    for task in plan['tasks']:
        src_label = next(label for label in task['labels'] if label.startswith('src-sb-row-'))
        if src_label in existing:
            issue_key = existing[src_label]
            created.append({'kind': 'existing', 'key': issue_key, 'source_id': task['source_id']})
            continue
        issue_key = create_task(client, project_key, task)
        transition_if_needed(client, issue_key, task.get('status', ''))
        created.append({'kind': 'task', 'key': issue_key, 'source_id': task['source_id']})
    return {
        'project_key': project_key,
        'created': created,
        'created_count': len([x for x in created if x.get('kind') == 'task']),
    }


def main():
    parser = argparse.ArgumentParser(description='Import only Sprint Backlog rows as board items')
    add_env_arg(parser)
    parser.add_argument('--file', required=True)
    parser.add_argument('--project', required=True)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    plan = build_plan(args.file, args.project)
    if not args.apply:
        print_json(plan)
        return
    client = JiraClient(args.env_file)
    result = apply_plan(client, plan)
    print_json({**result, 'count': plan['count']})


if __name__ == '__main__':
    main()
