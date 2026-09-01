#!/usr/bin/env python3
import argparse
import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from jira_common import JiraClient, adf_text, add_env_arg, print_json

ALIASES = {
    "issue_key": ["issue_key", "key", "jira_key", "ticket", "ticket_key", "issue"],
    "summary": ["summary", "title", "task", "issue_summary"],
    "description": ["description", "desc", "details"],
    "comment": ["comment", "note", "update", "remarks"],
    "status": ["status", "transition", "state"],
    "priority": ["priority"],
    "labels": ["labels", "tags"],
}


def normalize(name: str) -> str:
    return name.strip().lower().replace(" ", "_")


def resolve_columns(headers: List[str]) -> Dict[str, Optional[str]]:
    norm_map = {normalize(h): h for h in headers if h is not None}
    resolved: Dict[str, Optional[str]] = {}
    for canonical, aliases in ALIASES.items():
        resolved[canonical] = None
        for alias in aliases:
            if alias in norm_map:
                resolved[canonical] = norm_map[alias]
                break
    return resolved


def load_csv(path: Path) -> Tuple[List[str], List[Dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        rows = list(reader)
        return list(reader.fieldnames or []), rows


def load_xlsx(path: Path, sheet: Optional[str]) -> Tuple[List[str], List[Dict[str, Any]]]:
    wb = load_workbook(filename=path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return [], []
    headers = [str(v).strip() if v is not None else "" for v in data[0]]
    rows = []
    for row in data[1:]:
        if row is None:
            continue
        item = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            value = row[idx] if idx < len(row) else None
            if value is None:
                continue
            item[header] = str(value).strip() if not isinstance(value, str) else value.strip()
        if item:
            rows.append(item)
    return headers, rows


def load_rows(file_path: str, sheet: Optional[str]) -> Tuple[List[str], List[Dict[str, Any]]]:
    path = Path(file_path).expanduser()
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return load_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return load_xlsx(path, sheet)
    raise SystemExit(f"Unsupported file type: {suffix}")


def get_value(row: Dict[str, Any], column: Optional[str]) -> str:
    if not column:
        return ""
    value = row.get(column, "")
    return value.strip() if isinstance(value, str) else str(value).strip()


def parse_labels(value: str) -> List[str]:
    if not value:
        return []
    parts = [x.strip() for chunk in value.split(",") for x in chunk.split(";")]
    return [x for x in parts if x]


def build_plan(row_idx: int, row: Dict[str, Any], columns: Dict[str, Optional[str]], project: Optional[str]) -> Dict[str, Any]:
    issue_key = get_value(row, columns.get("issue_key"))
    summary = get_value(row, columns.get("summary"))
    description = get_value(row, columns.get("description"))
    comment = get_value(row, columns.get("comment"))
    status = get_value(row, columns.get("status"))
    priority = get_value(row, columns.get("priority"))
    labels = parse_labels(get_value(row, columns.get("labels")))

    plan = {
        "row": row_idx,
        "issue_key": issue_key,
        "project": project,
        "actions": [],
        "skipped": False,
        "reason": None,
    }

    if not issue_key:
        plan["skipped"] = True
        plan["reason"] = "missing issue key"
        return plan

    fields: Dict[str, Any] = {}
    if summary:
        fields["summary"] = summary
    if description:
        fields["description"] = adf_text(description)
    if priority:
        fields["priority"] = {"name": priority}
    if labels:
        fields["labels"] = labels
    if fields:
        plan["actions"].append({"type": "update_fields", "fields": fields})
    if comment:
        plan["actions"].append({"type": "comment", "body": comment})
    if status:
        plan["actions"].append({"type": "transition", "name": status})
    if not plan["actions"]:
        plan["skipped"] = True
        plan["reason"] = "no supported update columns populated"
    return plan


def transition_id_by_name(client: JiraClient, issue_key: str, name: str) -> str:
    wanted = name.strip().lower()
    transitions = client.transitions(issue_key).get("transitions", [])
    matched = next((t for t in transitions if t.get("name", "").strip().lower() == wanted), None)
    if not matched:
        available = [t.get("name") for t in transitions]
        raise RuntimeError(f"No transition named '{name}' for {issue_key}. Available: {available}")
    return str(matched["id"])


def apply_plan(client: JiraClient, plan: Dict[str, Any]) -> Dict[str, Any]:
    results = []
    for action in plan["actions"]:
        if action["type"] == "update_fields":
            results.append({"type": "update_fields", "result": client.issue_update(plan["issue_key"], action["fields"])})
        elif action["type"] == "comment":
            results.append({"type": "comment", "result": client.issue_comment(plan["issue_key"], action["body"])})
        elif action["type"] == "transition":
            transition_id = transition_id_by_name(client, plan["issue_key"], action["name"])
            results.append({"type": "transition", "transition_id": transition_id, "result": client.transition(plan["issue_key"], transition_id)})
    return {
        "row": plan["row"],
        "issue_key": plan["issue_key"],
        "results": results,
    }


def main():
    parser = argparse.ArgumentParser(description="Preview or apply Jira updates from CSV/XLSX")
    add_env_arg(parser)
    parser.add_argument("--file", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--project")
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    headers, rows = load_rows(args.file, args.sheet)
    columns = resolve_columns(headers)
    plans = [build_plan(idx + 2, row, columns, args.project) for idx, row in enumerate(rows)]

    summary = {
        "file": str(Path(args.file).expanduser()),
        "sheet": args.sheet,
        "headers": headers,
        "resolved_columns": columns,
        "row_count": len(rows),
        "planned": len([p for p in plans if not p["skipped"]]),
        "skipped": len([p for p in plans if p["skipped"]]),
        "apply": args.apply,
        "plans": plans,
    }

    if not args.apply:
        print_json(summary)
        return

    client = JiraClient(args.env_file)
    results = []
    for plan in plans:
        if plan["skipped"]:
            continue
        results.append(apply_plan(client, plan))

    print_json({
        **summary,
        "results": results,
    })


if __name__ == "__main__":
    main()
