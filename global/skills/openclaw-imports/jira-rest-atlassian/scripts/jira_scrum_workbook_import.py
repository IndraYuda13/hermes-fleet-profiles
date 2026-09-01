#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from jira_common import JiraClient, adf_text, add_env_arg, print_json

STOPWORDS = {
    'sebagai','saya','ingin','agar','dan','yang','di','ke','dengan','untuk','dapat','data','melalui',
    'pada','agar','akun','aplikasi','mobile','web','admin','user','pengguna','dapat','agar','dari','via',
    'the','a','an','of','to','in','on'
}


def slug(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r'[^a-z0-9]+', '-', value)
    return value.strip('-') or 'item'


def tokens(text: str) -> set:
    parts = re.split(r'[^a-zA-Z0-9]+', text.lower())
    return {p for p in parts if p and p not in STOPWORDS and len(p) > 1}


def priority_name(raw: str) -> Optional[str]:
    if not raw:
        return None
    norm = raw.strip().lower()
    mapping = {
        'highest': 'Highest',
        'high': 'High',
        'medium': 'Medium',
        'low': 'Low',
        'lowest': 'Lowest',
    }
    return mapping.get(norm)


def sprint_label(name: str) -> Optional[str]:
    if not name:
        return None
    return slug(name)


def clean(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip()


def load_workbook_rows(path: str):
    wb = load_workbook(filename=Path(path).expanduser(), data_only=True)
    def sheet_rows(sheet_name: str):
        ws = wb[sheet_name]
        headers = [clean(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            item = {headers[idx]: clean(row[idx]) if idx < len(headers) else '' for idx in range(len(headers)) if headers[idx]}
            rows.append(item)
        return rows
    return {
        'user_stories': sheet_rows('User Stories'),
        'product_backlog': sheet_rows('Product Backlog'),
        'sprint_backlog': sheet_rows('Sprint Backlog'),
    }


def module_for_pbi(pbi_text: str, user_stories: List[Dict[str, str]]) -> str:
    pbi_tokens = tokens(pbi_text)
    best_module = 'General'
    best_score = -1
    for row in user_stories:
        story_text = row.get('User Story (format Persona-Need-Goal)', '')
        module = row.get('Epic/Modul', 'General') or 'General'
        score = len(pbi_tokens & tokens(story_text))
        if score > best_score:
            best_score = score
            best_module = module
    return best_module


def build_description(lines: List[Tuple[str, str]]) -> Dict[str, Any]:
    chunks = []
    for title, value in lines:
        if not value:
            continue
        chunks.append(f"{title}: {value}")
    return adf_text('\n'.join(chunks))


def build_plan(file_path: str, project_key: str, sprint_mode: str) -> Dict[str, Any]:
    data = load_workbook_rows(file_path)
    user_stories = data['user_stories']
    product_backlog = data['product_backlog']
    sprint_backlog = data['sprint_backlog']

    epics: Dict[str, Dict[str, Any]] = {}
    for row in user_stories:
        module = row.get('Epic/Modul', 'General') or 'General'
        if module not in epics:
            epics[module] = {
                'type': 'Epic',
                'source_id': f'EPIC-{slug(module)}',
                'summary': module,
                'labels': ['import-simasu', f"src-epic-{slug(module)}"],
                'description': build_description([
                    ('Source', 'Workbook User Stories'),
                    ('Module', module),
                ]),
            }

    stories: List[Dict[str, Any]] = []
    for row in user_stories:
        module = row.get('Epic/Modul', 'General') or 'General'
        us_id = row.get('ID', '')
        summary = row.get('User Story (format Persona-Need-Goal)', '')
        pri = priority_name(row.get('Prioritas', ''))
        sp = row.get('Estimasi (SP)', '')
        ac = row.get('Acceptance Criteria', '')
        labels = ['import-simasu', 'user-story', f"src-{slug(us_id)}", f"module-{slug(module)}"]
        stories.append({
            'type': 'Story',
            'source_id': us_id,
            'summary': summary,
            'priority': pri,
            'labels': labels,
            'module': module,
            'parent_ref': f'EPIC-{slug(module)}',
            'description': build_description([
                ('Source ID', us_id),
                ('Module', module),
                ('Story Points', sp),
                ('Acceptance Criteria', ac),
            ]),
        })

    features: List[Dict[str, Any]] = []
    pbi_ref_to_plan: Dict[str, Dict[str, Any]] = {}
    for row in product_backlog:
        pbi_id = row.get('ID', '')
        summary = row.get('User Story / PBI', '')
        pri = priority_name(row.get('Prioritas', ''))
        bv = row.get('Nilai Bisnis (BV)', '')
        sp = row.get('Estimasi (SP)', '')
        sprint_target = row.get('Sprint Target', '')
        status = row.get('Status', '')
        ac = row.get('Acceptance Criteria (Ringkas)', '')
        module = module_for_pbi(summary, user_stories)
        labels = ['import-simasu', 'product-backlog', f"src-{slug(pbi_id)}", f"module-{slug(module)}"]
        if sprint_mode == 'labels' and sprint_target:
            labels.append(f"sprint-{sprint_label(sprint_target)}")
        item = {
            'type': 'Feature',
            'source_id': pbi_id,
            'summary': summary,
            'priority': pri,
            'status': status,
            'labels': labels,
            'module': module,
            'parent_ref': f'EPIC-{slug(module)}',
            'description': build_description([
                ('Source ID', pbi_id),
                ('Module', module),
                ('Business Value', bv),
                ('Story Points', sp),
                ('Sprint Target', sprint_target),
                ('Acceptance Criteria', ac),
            ]),
        }
        features.append(item)
        pbi_ref_to_plan[pbi_id] = item

    subtasks: List[Dict[str, Any]] = []
    for idx, row in enumerate(sprint_backlog, start=2):
        sprint = row.get('Sprint', '')
        sprint_goal = row.get('Sprint Goal', '')
        pbi_id = row.get('PBI ID', '')
        task = row.get('Task', '')
        pic = row.get('PIC', '')
        est = row.get('Estimasi (Jam)', '')
        status = row.get('Status', '')
        dod = row.get('Catatan/DoD', '')
        labels = ['import-simasu', 'sprint-task', f'src-sb-row-{idx}', f"src-{slug(pbi_id)}"]
        if sprint_mode == 'labels' and sprint:
            labels.append(f"sprint-{sprint_label(sprint)}")
        subtasks.append({
            'type': 'Subtask',
            'source_id': f'SB-ROW-{idx}',
            'summary': task,
            'status': status,
            'labels': labels,
            'parent_ref': pbi_id,
            'description': build_description([
                ('Source Row', str(idx)),
                ('Sprint', sprint),
                ('Sprint Goal', sprint_goal),
                ('PBI ID', pbi_id),
                ('PIC', pic),
                ('Estimate Hours', est),
                ('Definition of Done', dod),
            ]),
        })

    return {
        'project_key': project_key,
        'sprint_mode': sprint_mode,
        'counts': {
            'epics': len(epics),
            'user_stories': len(stories),
            'product_backlog_items': len(features),
            'sprint_tasks': len(subtasks),
            'total_issues_to_create': len(epics) + len(stories) + len(features) + len(subtasks),
        },
        'epics': list(epics.values()),
        'stories': stories,
        'features': features,
        'subtasks': subtasks,
    }


def existing_imported(client: JiraClient, project_key: str) -> Dict[str, str]:
    result = client.search(f'project = {project_key} AND labels = "import-simasu"', limit=200, fields='labels,summary')
    mapping = {}
    for issue in result.get('issues', []):
        labels = issue.get('fields', {}).get('labels', []) or []
        for label in labels:
            if label.startswith('src-'):
                mapping[label] = issue['key']
    return mapping


def project_issue_types(client: JiraClient, project_key: str) -> set:
    data = client.project(project_key)
    return {item.get('name') for item in data.get('issueTypes', []) if item.get('name')}


def resolve_issue_type(requested: str, available: set) -> str:
    if requested in available:
        return requested
    fallback_map = {
        'Feature': 'Task',
        'Story': 'Task',
        'Epic': 'Epic',
        'Subtask': 'Subtask',
    }
    fallback = fallback_map.get(requested)
    if fallback and fallback in available:
        return fallback
    raise RuntimeError(f"No compatible issue type for {requested}. Available: {sorted(available)}")


def create_issue(client: JiraClient, project_key: str, available_types: set, issue_type: str, summary: str, description: Dict[str, Any], labels: List[str], priority: Optional[str] = None, parent_key: Optional[str] = None) -> str:
    resolved_type = resolve_issue_type(issue_type, available_types)
    fields: Dict[str, Any] = {
        'project': {'key': project_key},
        'issuetype': {'name': resolved_type},
        'summary': summary,
        'description': description,
        'labels': labels,
    }
    if priority and resolved_type != 'Epic':
        fields['priority'] = {'name': priority}
    if parent_key:
        fields['parent'] = {'key': parent_key}
    result = client.issue_create({'fields': fields})
    return result['key']


def transition_if_needed(client: JiraClient, issue_key: str, desired_status: str):
    if not desired_status or desired_status.strip().lower() in {'to do', 'todo', 'backlog'}:
        return
    transitions = client.transitions(issue_key).get('transitions', [])
    desired = desired_status.strip().lower()
    match = next((t for t in transitions if t.get('name', '').strip().lower() == desired), None)
    if match:
        client.transition(issue_key, match['id'])


def apply_plan(client: JiraClient, plan: Dict[str, Any]) -> Dict[str, Any]:
    project_key = plan['project_key']
    existing = existing_imported(client, project_key)
    created = []
    available_types = project_issue_types(client, project_key)

    # Epics first
    epic_key_by_ref: Dict[str, str] = {}
    for epic in plan['epics']:
        src_label = next(label for label in epic['labels'] if label.startswith('src-'))
        if src_label in existing:
            epic_key = existing[src_label]
        else:
            epic_key = create_issue(client, project_key, available_types, epic['type'], epic['summary'], epic['description'], epic['labels'])
            created.append({'kind': 'epic', 'key': epic_key, 'source_id': epic['source_id']})
        epic_key_by_ref[epic['source_id']] = epic_key

    # Features/PBIs next
    feature_key_by_pbi: Dict[str, str] = {}
    for feature in plan['features']:
        src_label = next(label for label in feature['labels'] if label.startswith('src-pbi-'))
        parent_key = epic_key_by_ref[feature['parent_ref']]
        if src_label in existing:
            issue_key = existing[src_label]
        else:
            issue_key = create_issue(client, project_key, available_types, feature['type'], feature['summary'], feature['description'], feature['labels'], feature.get('priority'), parent_key)
            transition_if_needed(client, issue_key, feature.get('status', ''))
            created.append({'kind': 'feature', 'key': issue_key, 'source_id': feature['source_id']})
        feature_key_by_pbi[feature['source_id']] = issue_key

    # Stories
    for story in plan['stories']:
        src_label = next(label for label in story['labels'] if label.startswith('src-us-'))
        parent_key = epic_key_by_ref[story['parent_ref']]
        if src_label in existing:
            issue_key = existing[src_label]
        else:
            issue_key = create_issue(client, project_key, available_types, story['type'], story['summary'], story['description'], story['labels'], story.get('priority'), parent_key)
            created.append({'kind': 'story', 'key': issue_key, 'source_id': story['source_id']})

    # Sprint tasks as subtasks under PBI feature/task
    for subtask in plan['subtasks']:
        src_label = next(label for label in subtask['labels'] if label.startswith('src-sb-row-'))
        parent_key = feature_key_by_pbi.get(subtask['parent_ref'])
        if not parent_key:
            created.append({'kind': 'subtask-skip', 'source_id': subtask['source_id'], 'reason': f"missing parent PBI {subtask['parent_ref']}"})
            continue
        if src_label in existing:
            issue_key = existing[src_label]
        else:
            issue_key = create_issue(client, project_key, available_types, subtask['type'], subtask['summary'], subtask['description'], subtask['labels'], None, parent_key)
            transition_if_needed(client, issue_key, subtask.get('status', ''))
            created.append({'kind': 'subtask', 'key': issue_key, 'source_id': subtask['source_id']})

    return {
        'project_key': project_key,
        'available_types': sorted(available_types),
        'created': created,
        'created_count': len([x for x in created if 'key' in x]),
    }


def main():
    parser = argparse.ArgumentParser(description='Preview or apply SIMASU scrum workbook import into Jira')
    add_env_arg(parser)
    parser.add_argument('--file', required=True)
    parser.add_argument('--project', required=True)
    parser.add_argument('--sprint-mode', choices=['labels', 'skip'], default='labels')
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()

    plan = build_plan(args.file, args.project, args.sprint_mode)
    if not args.apply:
        print_json(plan)
        return
    client = JiraClient(args.env_file)
    result = apply_plan(client, plan)
    print_json({**plan['counts'], **result, 'sprint_mode': args.sprint_mode})


if __name__ == '__main__':
    main()
